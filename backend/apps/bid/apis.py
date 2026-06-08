"""M4 标书撰写 — API"""
from fastapi import APIRouter, Depends, UploadFile, File
from pydantic import BaseModel
from apps.qa.services.llm_client import LlmClient, LlmRequest
from core.settings import get_settings
import re, time, base64, io
from docx import Document

router = APIRouter(prefix="/bid", tags=["标书撰写"])

TEMPLATE_DIR = "../Ai招投标模板/"

class FillRequest(BaseModel):
    industry: str
    project_info: str
    template_name: str = "商务标——模板(1).docx"

class FillResponse(BaseModel):
    sections: list[dict]
    total_tokens: int = 0
    total_latency_ms: int = 0

# 模板内存存储
_templates: dict[str, bytes] = {}

@router.get("/templates")
async def list_templates():
    return [
        {"name": "商务标——模板(1).docx", "label": "商务标模板", "size": "55KB"},
        {"name": "技术标——模板(1).docx", "label": "技术标模板", "size": "947KB"},
    ]

@router.post("/upload")
async def upload_template(file: UploadFile = File(...)):
    data = await file.read()
    _templates[file.filename] = data
    # 提取占位符
    placeholders = []
    try:
        doc = Document(io.BytesIO(data))
        for p in doc.paragraphs:
            for m in re.finditer(r"\{\{(.+?)}}", p.text):
                if m.group(1).strip() not in placeholders:
                    placeholders.append(m.group(1).strip())
    except: pass
    return {"name": file.filename, "placeholders": placeholders}

@router.post("/fill", response_model=FillResponse)
async def fill(req: FillRequest, settings=Depends(get_settings)):
    llm = LlmClient(settings.llm_api_key, settings.llm_base_url, settings.llm_model)
    start = time.time()

    # 读取模板提取章节标题
    template_path = TEMPLATE_DIR + req.template_name
    sections = _extract_sections(template_path)
    if not sections:
        sections = ["报价一览表", "报价明细表"]

    # AI 填充每节
    results = []
    total_tokens = 0
    for title in sections:
        prompt = (f"撰写商务标书【{title}】部分。\n行业:{req.industry}\n项目:{req.project_info}\n\n"
                  f"要求: 直接写正文，不写元话语。如有表格用Markdown格式。")
        try:
            r = await llm.call(LlmRequest(system_prompt="你是商务标书撰写专家。", user_message=prompt, max_tokens=1500, temperature=0.5))
            cleaned = re.sub(r"^(好的[，,]?\s*)?(作为.*?[,，]\s*)?(我将.*?[。\s]*)?", "", r.content).strip()
            results.append({"title": title, "content": cleaned})
            total_tokens += r.tokens_used
        except Exception as e:
            results.append({"title": title, "content": f"(生成失败: {e})"})

    return FillResponse(sections=results, total_tokens=total_tokens,
                        total_latency_ms=int((time.time()-start)*1000))

def _extract_sections(path: str) -> list[str]:
    try:
        doc = Document(path)
        titles = []
        for p in doc.paragraphs:
            s = p.style.name if p.style else ""
            t = p.text.strip()
            if ("heading" in s.lower() or "标题" in s) and t:
                titles.append(t)
        return titles if titles else ["报价一览表", "报价明细表"]
    except:
        return ["报价一览表", "报价明细表"]
